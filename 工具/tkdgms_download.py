# 下載 tkdgms 索引到的檔案，全文搜名字
import requests, os, sys, json, re
from urllib.parse import unquote
sys.stdout.reconfigure(encoding='utf-8')
OUT=r"C:\Users\TUF Gaming\Desktop\我的專案\林家兄弟跆拳道\紀錄\其他縣市\tkdgms"
idx=json.load(open(os.path.join(OUT,"_index.json"),encoding="utf-8"))
NAMES=["林聖翔","林聖宸","林聖辰","聖翔","聖宸","聖辰","民逸","福山"]
S=requests.Session(); S.headers["User-Agent"]="Mozilla/5.0"
def text_of(path):
    ext=path.lower().rsplit(".",1)[-1]
    try:
        if ext=="pdf":
            import pdfplumber
            with pdfplumber.open(path) as pdf: return "\n".join((p.extract_text() or "") for p in pdf.pages)
        if ext in("xlsx","xlsm"):
            import openpyxl; wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
            return "\n".join(" ".join(str(c) for c in row if c is not None) for ws in wb for row in ws.iter_rows(values_only=True))
        if ext=="xls":
            import xlrd; wb=xlrd.open_workbook(path); return "\n".join(" ".join(str(c) for c in ws.row_values(r)) for ws in wb.sheets() for r in range(ws.nrows))
        if ext=="docx":
            import docx; d=docx.Document(path); return "\n".join(p.text for p in d.paragraphs)+"\n"+"\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
    except Exception as e: return f"[ERR {e}]"
    return ""
report=[]
for u,meta in idx["files"].items():
    fn=unquote(u.rsplit("/",1)[-1]).split("?")[0]
    if not re.search(r'\.(pdf|xlsx?|docx?)$',fn,re.I): report.append({"url":u,"name":fn,"page":meta["page"],"note":"非檔案連結，跳過"}); continue
    path=os.path.join(OUT,fn)
    if not os.path.exists(path):
        try:
            r=S.get(u,timeout=30); open(path,"wb").write(r.content)
        except Exception as e: report.append({"url":u,"name":fn,"note":f"下載失敗 {e}"}); continue
    t=text_of(path)
    hits={n:t.count(n) for n in NAMES if n in t}
    report.append({"url":u,"name":fn,"page":meta["page"],"chars":len(t),"hits":hits})
    print(fn,"|",meta["page"][:30],"| chars",len(t),"| hits",hits)
json.dump(report,open(os.path.join(OUT,"_搜尋結果.json"),"w",encoding="utf-8"),ensure_ascii=False,indent=1)
